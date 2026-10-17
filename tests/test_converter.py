import os
import fitz
from pathlib import Path
from docx import Document
import pytest
from backend.app.services.pdf_to_docx import convert_pdf_to_docx
from backend.app.services.docx_to_pdf import convert_docx_to_pdf
from backend.app.services.pdf_compress import compress_pdf
from backend.app.services.pdf_ops import merge_pdfs, split_pdf
from backend.app.core.validator import validate_pdf_file, validate_docx_file
from backend.app.core.exceptions import ValidationError

TEST_DIR = Path(__file__).parent / "test_scratch"

@pytest.fixture(scope="session", autouse=True)
def setup_scratch():
    TEST_DIR.mkdir(parents=True, exist_ok=True)
    yield
    import shutil
    shutil.rmtree(TEST_DIR, ignore_errors=True)

def create_sample_pdf(path: Path, text: str = "Hello World! Test PDF Converter.") -> Path:
    """Generates a test PDF with text."""
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), text, fontsize=14)
    doc.save(str(path))
    doc.close()
    return path

def create_sample_docx(path: Path, text: str = "Sample DOCX Title") -> Path:
    """Generates a test DOCX."""
    doc = Document()
    doc.add_heading(text, level=1)
    doc.add_paragraph("This is a test paragraph for DOCX conversion.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Item"
    table.cell(0, 1).text = "Price"
    table.cell(1, 0).text = "Book"
    table.cell(1, 1).text = "$15"
    doc.save(str(path))
    return path

@pytest.mark.asyncio
async def test_pdf_to_docx_conversion():
    sample_pdf = create_sample_pdf(TEST_DIR / "sample.pdf")
    output_docx = TEST_DIR / "converted.docx"
    
    await convert_pdf_to_docx(sample_pdf, output_docx)
    
    assert output_docx.exists()
    assert output_docx.stat().st_size > 0
    # Verify it can be opened as a docx
    doc = Document(str(output_docx))
    all_text = " ".join([p.text for p in doc.paragraphs])
    assert "Hello World" in all_text

@pytest.mark.asyncio
async def test_docx_to_pdf_conversion():
    sample_docx = create_sample_docx(TEST_DIR / "sample.docx")
    
    output_pdf = await convert_docx_to_pdf(sample_docx, TEST_DIR)
    
    assert output_pdf.exists()
    assert output_pdf.stat().st_size > 0
    # Verify it is a valid PDF
    pdf_doc = fitz.open(output_pdf)
    assert len(pdf_doc) > 0
    pdf_doc.close()

@pytest.mark.asyncio
async def test_pdf_compression():
    sample_pdf = create_sample_pdf(TEST_DIR / "sample_comp.pdf")
    out_pdf = TEST_DIR / "compressed.pdf"
    
    path, orig_s, comp_s = await compress_pdf(sample_pdf, out_pdf)
    assert out_pdf.exists()
    assert comp_s > 0

@pytest.mark.asyncio
async def test_pdf_merge_and_split():
    pdf1 = create_sample_pdf(TEST_DIR / "part1.pdf", "Page 1 Content")
    pdf2 = create_sample_pdf(TEST_DIR / "part2.pdf", "Page 2 Content")
    merged = TEST_DIR / "merged.pdf"
    
    # Merge
    await merge_pdfs([pdf1, pdf2], merged)
    assert merged.exists()
    doc = fitz.open(merged)
    assert len(doc) == 2
    doc.close()

    # Split
    split_out = TEST_DIR / "page1_only.pdf"
    await split_pdf(merged, split_out, "1")
    assert split_out.exists()
    doc_split = fitz.open(split_out)
    assert len(doc_split) == 1
    doc_split.close()

def test_validation_empty_file():
    empty_file = TEST_DIR / "empty.pdf"
    empty_file.touch()
    with pytest.raises(ValidationError):
        validate_pdf_file(empty_file)
